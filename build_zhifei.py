import subprocess
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

WB = '/workspace/产品成本测算报价表.xlsx'
RECALC = '/data/user/plugins/trae-remote-official/office-plugin/0.1.0/skills/xlsx/scripts/recalc.py'

F_NAME = '微软雅黑'
BLUE = 'FF0000FF'
BLACK = 'FF000000'
GREEN = 'FF008000'
WHITE = 'FFFFFFFF'
FILL_TITLE = PatternFill('solid', fgColor='FF0F3B5D')
FILL_SEC = PatternFill('solid', fgColor='FF3288B9')
FILL_KPI = PatternFill('solid', fgColor='FFEAF2FF')
THIN = Side(style='thin', color='FFD9DEE7')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(size=10, bold=False, color=BLACK):
    return Font(name=F_NAME, size=size, bold=bold, color=color)


wb = openpyxl.load_workbook(WB)

idx = None
for old in ('加工及工装', '制费明细'):
    if old in wb.sheetnames:
        idx = wb.sheetnames.index(old)
        del wb[old]
        break
if idx is None:
    idx = len(wb.sheetnames)
ws = wb.create_sheet('制费明细', idx)

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 14
ws.column_dimensions['C'].width = 22
ws.column_dimensions['D'].width = 13
ws.column_dimensions['E'].width = 11
ws.column_dimensions['F'].width = 11
ws.column_dimensions['G'].width = 13
ws.column_dimensions['H'].width = 40


def put(cell, value, f=None, fill=None, fmt=None, align=None, border=True):
    c = ws[cell]
    c.value = value
    c.font = f or font()
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = align or Alignment(vertical='center', wrap_text=True)
    if border:
        c.border = BORDER
    return c


def title_row(row, text, span='A:H'):
    ws.merge_cells(f'A{row}:H{row}')
    put(f'A{row}', text, font(11, True, WHITE), FILL_SEC, border=False)


def header_row(row, headers):
    for col, h in zip('ABCDEFGH', headers):
        put(f'{col}{row}', h, font(10, True), FILL_KPI,
            align=Alignment(horizontal='center', vertical='center', wrap_text=True))


FMT_MONEY = '#,##0.0000'
FMT_PRICE = '#,##0.000000'

# ---- 标题 ----
ws.merge_cells('A1:H1')
put('A1', '制费明细（工程报价表格式 P01~P07：加工费 + 辅料费 + 工装治具费）',
    font(14, True, WHITE), FILL_TITLE, border=False)
ws.row_dimensions[1].height = 26

# ---- 项目信息 ----
put('A3', '客户', font(10, True), FILL_KPI)
put('B3', '=报价汇总!B6', font(color=GREEN))
put('C3', '机型', font(10, True), FILL_KPI)
put('D3', '=报价汇总!B4', font(color=GREEN))
put('E3', '预计订单(pcs)', font(10, True), FILL_KPI)
put('F3', None, font())
ws.merge_cells('E3:F3')
put('G3', 100000, font(color=BLUE), fmt='#,##0')
put('H3', '夹具一次性费用的分摊分母，按项目年需求量/合同量填写', font(9, color='FF808080'), border=False)

put('A4', '人工费率(元/秒)', font(10, True), FILL_KPI)
put('B4', 0.0089138889, font(color=BLUE), fmt=FMT_PRICE)
put('C4', '32.09元/小时÷3600', font(9, color='FF808080'), border=False)
put('D4', '宽放系数', font(10, True), FILL_KPI)
put('E4', 1.1, font(color=BLUE), fmt='0.00')
put('F4', '烧录/成型/组装共用', font(9, color='FF808080'), border=False)
ws.merge_cells('G4:H4')
put('G4', '费率来源：003/009参考报价反推（0.0588元÷6.6s），可按实际改', font(9, color='FF808080'), border=False)

RATE = '$B$4'
KF = '$E$4'

# ---- P01 SMT ----
title_row(6, 'P01  SMT贴片（费用 = 工艺单价 × 元件个数）')
header_row(7, ['项目', '工序名称', '元件类型', '工艺单价(元)', '元件个数', '备注', '费用(元)', '说明'])
smt = [
    ('0603/0402', 0.019, 56, ''),
    ('1206/0805', 0.02, 4, ''),
    ('二/三极管', 0.02, 0, ''),
    ('按键/插座', 0.06, 4, ''),
    ('电感/MOS管/EC等', 0.035, 1, ''),
    ('贴片IC', 0.018, 9, '每4个pin为1个元件'),
    ('模块', 0.142, 0, ''),
]
r = 8
for name, price, cnt, note in smt:
    put(f'A{r}', 'P01' if r == 8 else None)
    put(f'B{r}', 'SMT' if r == 8 else None)
    put(f'C{r}', name)
    put(f'D{r}', price, font(color=BLUE), fmt=FMT_PRICE)
    put(f'E{r}', cnt, font(color=BLUE), fmt='#,##0')
    put(f'F{r}', note, font(9, color='FF808080'))
    put(f'G{r}', f'=IFERROR(D{r}*E{r},0)', fmt=FMT_MONEY)
    put(f'H{r}', None)
    r += 1
SMT_FIRST, SMT_LAST = 8, r - 1
put(f'C{r}', 'SMT成本(元)', font(10, True), FILL_KPI)
put(f'G{r}', f'=SUM(G{SMT_FIRST}:G{SMT_LAST})', font(10, True), FILL_KPI, FMT_MONEY)
SMT_TOTAL = f'$G${r}'
r += 2

# ---- P02 烧录 ----
title_row(r, 'P02  烧录（工序工时 = 工作时间 × 工序个数 × 宽放；成本 = 工时 × 人工费率）')
r += 1
header_row(r, ['项目', '工序名称', '工艺流程', '工作时间(s)', '工序个数', '宽放系数', '工序工时(s)', '费用(元)'])
r += 1
sl = [('离线烧录', 6, 1), ('在线烧录', 3, 0)]
SL_FIRST = r
for name, t, cnt in sl:
    put(f'A{r}', 'P02' if r == SL_FIRST else None)
    put(f'B{r}', 'SL' if r == SL_FIRST else None)
    put(f'C{r}', name)
    put(f'D{r}', t, font(color=BLUE), fmt='0.0')
    put(f'E{r}', cnt, font(color=BLUE), fmt='#,##0')
    put(f'F{r}', f'={KF}', fmt='0.00')
    put(f'G{r}', f'=D{r}*E{r}*F{r}', fmt='0.00')
    put(f'H{r}', f'=G{r}*{RATE}', fmt=FMT_MONEY)
    r += 1
SL_LAST = r - 1
put(f'C{r}', '烧录成本(元)', font(10, True), FILL_KPI)
put(f'H{r}', f'=SUM(H{SL_FIRST}:H{SL_LAST})', font(10, True), FILL_KPI, FMT_MONEY)
SL_TOTAL = f'$H${r}'
r += 2

# ---- P03 AI/RI ----
title_row(r, 'P03  AI/RI 机插（费用 = 单价 × 个数 × 宽放）')
r += 1
header_row(r, ['项目', '工序名称', '元件类型', '工艺单价(元)', '元件个数', '宽放系数', '费用(元)', '说明'])
r += 1
ai = [('电阻/二极管', 0.025, 0), ('跳线', 0.022, 0), ('RI', 0.025, 0)]
AI_FIRST = r
for name, price, cnt in ai:
    put(f'A{r}', 'P03' if r == AI_FIRST else None)
    put(f'B{r}', 'AI/RI' if r == AI_FIRST else None)
    put(f'C{r}', name)
    put(f'D{r}', price, font(color=BLUE), fmt=FMT_PRICE)
    put(f'E{r}', cnt, font(color=BLUE), fmt='#,##0')
    put(f'F{r}', f'={KF}', fmt='0.00')
    put(f'G{r}', f'=IFERROR(D{r}*E{r}*F{r},0)', fmt=FMT_MONEY)
    put(f'H{r}', None)
    r += 1
AI_LAST = r - 1
put(f'C{r}', 'AI成本(元)', font(10, True), FILL_KPI)
put(f'G{r}', f'=SUM(G{AI_FIRST}:G{AI_LAST})', font(10, True), FILL_KPI, FMT_MONEY)
AI_TOTAL = f'$G${r}'
r += 2

# ---- P04 成型 ----
title_row(r, 'P04  成型（工序工时 = 工作时间 × 元件数量 × 宽放；成本 = 工时 × 人工费率）')
r += 1
header_row(r, ['项目', '工序名称', '工艺流程', '工作时间(s)', '元件数量', '宽放系数', '工序工时(s)', '费用(元)'])
r += 1
cx = [('电解电容', 1.2, 0), ('贴板成型', 1.8, 0), ('自动成型', 0.6, 0), ('晶闸管成型', 6, 0),
      ('散热器组装', 20, 0), ('光敏电阻成型', 8, 0), ('铆钉', 0.5, 0)]
CX_FIRST = r
for name, t, cnt in cx:
    put(f'A{r}', 'P04' if r == CX_FIRST else None)
    put(f'B{r}', '成型' if r == CX_FIRST else None)
    put(f'C{r}', name)
    put(f'D{r}', t, font(color=BLUE), fmt='0.0')
    put(f'E{r}', cnt, font(color=BLUE), fmt='#,##0')
    put(f'F{r}', f'={KF}', fmt='0.00')
    put(f'G{r}', f'=D{r}*E{r}*F{r}', fmt='0.00')
    put(f'H{r}', f'=G{r}*{RATE}', fmt=FMT_MONEY)
    r += 1
CX_LAST = r - 1
put(f'C{r}', '成型成本(元)', font(10, True), FILL_KPI)
put(f'H{r}', f'=SUM(H{CX_FIRST}:H{CX_LAST})', font(10, True), FILL_KPI, FMT_MONEY)
CX_TOTAL = f'$H${r}'
r += 2

# ---- P05 组装 ----
title_row(r, 'P05  组装（18道工序：工序工时 = 工作节拍 × 工序个数 × 宽放；成本 = 工时 × 人工费率）')
r += 1
header_row(r, ['项目', '工序名称', '工艺流程', '工作节拍(s)', '工序个数', '宽放系数', '工序工时(s)', '费用(元)'])
r += 1
zz = [('安装过炉治具', 3, 1), ('插件', 3, 1), ('插件QC', 3, 1), ('卸送治具/AVI检测', 3, 2),
      ('压件/执锡', 3, 2), ('炉后QC/洗板', 3, 2), ('ICT', 3, 1), ('分板', 3, 1),
      ('FCT', 3, 4), ('点胶', 3, 0), ('打螺钉', 3, 0), ('喷漆/收板/补胶', 3, 4),
      ('喷漆检查', 3, 1), ('贴标签', 3, 1), ('组装', 3, 0), ('FQC外观检查', 3, 2),
      ('包装', 3, 2), ('助拉', 3, 1)]
ZZ_FIRST = r
for name, t, cnt in zz:
    put(f'A{r}', 'P05' if r == ZZ_FIRST else None)
    put(f'B{r}', '组装' if r == ZZ_FIRST else None)
    put(f'C{r}', name)
    put(f'D{r}', t, font(color=BLUE), fmt='0.0')
    put(f'E{r}', cnt, font(color=BLUE), fmt='#,##0')
    put(f'F{r}', f'={KF}', fmt='0.00')
    put(f'G{r}', f'=D{r}*E{r}*F{r}', fmt='0.00')
    put(f'H{r}', f'=G{r}*{RATE}', fmt=FMT_MONEY)
    r += 1
ZZ_LAST = r - 1
put(f'C{r}', '组装成本(元)', font(10, True), FILL_KPI)
put(f'H{r}', f'=SUM(H{ZZ_FIRST}:H{ZZ_LAST})', font(10, True), FILL_KPI, FMT_MONEY)
ZZ_TOTAL = f'$H${r}'
r += 1

put(f'A{r}', '加工费合计(元)', font(11, True), FILL_KPI)
for col in 'BCDEFG':
    put(f'{col}{r}', None, fill=FILL_KPI)
ws.merge_cells(f'A{r}:F{r}')
put(f'H{r}', f'={SMT_TOTAL}+{SL_TOTAL}+{AI_TOTAL}+{CX_TOTAL}+{ZZ_TOTAL}',
    font(11, True), FILL_KPI, FMT_MONEY)
JGF_TOTAL = f'$H${r}'
r += 2

# ---- P06 辅料 ----
title_row(r, 'P06  辅料费用（费用 = 单价 × 单片用量 × 宽放；若辅料已计入BOM物料明细-辅助材料，请将用量置0避免重复）')
r += 1
header_row(r, ['项目', '费用类型', '辅料类型', '单价(元/g)', '单片用量(g)', '宽放系数', '辅料费用(元)', '说明'])
r += 1
fl = [('标签', 0.04, 2, 1.0, '元/个'), ('助焊剂', 0.019, 0.3, 1.2, ''), ('锡条(g)', 0.279, 0.1, 1.1, ''),
      ('锡丝(g)', 0.196, 0.1, 1.2, ''), ('胶类(g)', 0.04, 0, 1.1, ''), ('漆(g)', 0.21, 0.5, 1.2, '三防漆'),
      ('A/B胶(g)', 0.024, 0, 1.2, '')]
FL_FIRST = r
for name, price, use, kf, note in fl:
    put(f'A{r}', 'P06' if r == FL_FIRST else None)
    put(f'B{r}', '辅料费用' if r == FL_FIRST else None)
    put(f'C{r}', name)
    put(f'D{r}', price, font(color=BLUE), fmt=FMT_PRICE)
    put(f'E{r}', use, font(color=BLUE), fmt='0.000')
    put(f'F{r}', kf, font(color=BLUE), fmt='0.00')
    put(f'G{r}', f'=IFERROR(D{r}*E{r}*F{r},0)', fmt=FMT_MONEY)
    put(f'H{r}', note, font(9, color='FF808080'))
    r += 1
FL_LAST = r - 1
put(f'C{r}', '辅料成本(元)', font(10, True), FILL_KPI)
put(f'G{r}', f'=SUM(G{FL_FIRST}:G{FL_LAST})', font(10, True), FILL_KPI, FMT_MONEY)
FL_TOTAL = f'$G${r}'
r += 2

# ---- P07 夹具 ----
title_row(r, 'P07  夹具费用（不含样机阶段费用；单片分担 = 一次性总费用 ÷ 预计订单数）')
r += 1
header_row(r, ['项目', '费用类型', '夹具类型', '单价(元)', '个/张数', '备注', '夹具费用(元)', '说明'])
r += 1
jj = [('钢网', 265, 1), ('印刷辅助治具', 1000, 1), ('过锡炉治具', 350, 50), ('焊接治具', 1500, 0),
      ('铆钉模具', 20000, 0), ('分板治具', 800, 1), ('ICT治具', 8000, 1), ('烧录工装', 4000, 1),
      ('FCT工装', 8000, 4), ('OQC检验工装', 4000, 1), ('喷漆治具', 800, 20), ('灌胶治具', 300, 0)]
JJ_FIRST = r
for name, price, cnt in jj:
    put(f'A{r}', 'P07' if r == JJ_FIRST else None)
    put(f'B{r}', '夹具费用' if r == JJ_FIRST else None)
    put(f'C{r}', name)
    put(f'D{r}', price, font(color=BLUE), fmt='#,##0')
    put(f'E{r}', cnt, font(color=BLUE), fmt='#,##0')
    put(f'F{r}', None)
    put(f'G{r}', f'=IFERROR(D{r}*E{r},0)', fmt='#,##0.00')
    put(f'H{r}', None)
    r += 1
JJ_LAST = r - 1
put(f'C{r}', '夹具一次性总费用(元)', font(10, True), FILL_KPI)
put(f'G{r}', f'=SUM(G{JJ_FIRST}:G{JJ_LAST})', font(10, True), FILL_KPI, '#,##0.00')
JJ_SUM = f'$G${r}'
r += 1
put(f'C{r}', '夹具单片分担费用(元)', font(10, True), FILL_KPI)
put(f'G{r}', f'=IFERROR({JJ_SUM}/$G$3,0)', font(10, True), FILL_KPI, FMT_MONEY)
put(f'H{r}', '单片分担=一次性总费用÷预计订单数(G3)', font(9, color='FF808080'))
JJ_UNIT = f'$G${r}'
r += 1

put(f'A{r}', '单片制造成本合计(元)', font(11, True), FILL_KPI)
for col in 'BCDEFG':
    put(f'{col}{r}', None, fill=FILL_KPI)
ws.merge_cells(f'A{r}:G{r}')
put(f'H{r}', f'={JGF_TOTAL}+{FL_TOTAL}+{JJ_UNIT}', font(11, True), FILL_KPI, FMT_MONEY)
r += 2
put(f'A{r}', '制作：             审核：             核准：             批准日期：', font(9), border=False)

# ---- 报价汇总 重接 ----
wz = wb['报价汇总']
wz['A33'] = '加工费'
wz['B33'] = f'=制费明细!{JGF_TOTAL}'
wz['D33'] = '制费明细 P01~P05（SMT/烧录/AI/成型/组装，已含装配人工）'
wz['A34'] = '辅料费'
wz['B34'] = f'=制费明细!{FL_TOTAL}'
wz['D34'] = '制费明细 P06；若辅料已进BOM(辅助材料)请在P06置0，避免重复'
wz['A35'] = '工装分摊'
wz['B35'] = f'=制费明细!{JJ_UNIT}'
wz['D35'] = '制费明细 P07 夹具一次性总费用÷预计订单数'
for cell in ('B33', 'B34', 'B35'):
    wz[cell].font = font(10, color=GREEN)

# ---- 使用说明 追加 ----
wy = wb['使用说明']
for row in wy.iter_rows():
    for c in row:
        if isinstance(c.value, str) and '加工及工装' in c.value:
            c.value = c.value.replace('加工及工装', '制费明细')
existing = ' '.join(str(c.value) for row in wy.iter_rows() for c in row if c.value)
if '制费明细' not in existing:
    nr = wy.max_row + 1
    wy.cell(row=nr, column=1,
            value='制费明细：P01~P07工程报价表格式，蓝色单元格按项目填写（元件个数/工序个数/用量/张数/预计订单数），'
                  '默认预填003小板参考数据；辅料与BOM辅助材料二选一，避免重复计费。').font = font(10)

wb.save(WB)
print('saved. 制费明细 rows =', ws.max_row)
print('KEY CELLS:', 'JGF', JGF_TOTAL, 'FL', FL_TOTAL, 'JJ_UNIT', JJ_UNIT)

subprocess.run([sys.executable, RECALC, WB], check=True)
print('recalc done')
