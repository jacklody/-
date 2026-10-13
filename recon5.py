import openpyxl
wb = openpyxl.load_workbook('/workspace/产品成本测算报价表.xlsx', data_only=True)
ws = wb['跨越费率']
print('max_row', ws.max_row, 'max_col', ws.max_column)
for r in range(1, 8):
    print(r, [ws.cell(row=r, column=c).value for c in range(1, 18)])
print('--- col A distinct (产品) rows4-300')
import collections
ca = collections.Counter(ws.cell(row=r, column=1).value for r in range(4, min(ws.max_row, 300) + 1))
print(ca)
print('--- col Q samples')
qs = [(r, ws.cell(row=r, column=17).value) for r in range(4, min(ws.max_row, 300) + 1)]
print(qs[:6])
cities = collections.Counter()
for r, v in qs:
    if isinstance(v, str):
        for part in v.replace('，', ',').split(','):
            part = part.strip()
            if part:
                cities[part] += 1
print('n_cities', len(cities))
print(sorted(cities.items(), key=lambda x: -x[1])[:60])
wb2 = openpyxl.load_workbook('/workspace/产品成本测算报价表.xlsx')
w = wb2['物料明细']
print('--- 物料明细 H7,H110,I3 dv', repr(w['H7'].value), repr(w['H110'].value))
print('dv', [(str(dv.sqref), dv.formula1) for dv in w.data_validations.dataValidation])
print('row heights', {k: v.height for k, v in list(w.row_dimensions.items())[:5]})
print('autofilter', w.auto_filter.ref)
