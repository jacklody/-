import openpyxl
from openpyxl.utils import get_column_letter
wb = openpyxl.load_workbook('/workspace/产品成本测算报价表.xlsx')
out = []

ws = wb['报价汇总']
out.append('=== 报价汇总 rows 1-26 A-H ===')
for r in range(1, 27):
    vals = []
    for c in range(1, 9):
        v = ws.cell(r, c).value
        if v is not None:
            vals.append(f'{get_column_letter(c)}{r}={v!r}')
    if vals:
        out.append(' | '.join(vals))
out.append('merges: ' + repr([str(m) for m in ws.merged_cells.ranges]))
out.append('DVs:')
for dv in ws.data_validations.dataValidation:
    out.append(f'  {dv.sqref} type={dv.type} formula1={dv.formula1!r}')
out.append('freeze: ' + repr(ws.freeze_panes))

for name, cell in [('包装计算', 'B43'), ('运费计算', 'B33')]:
    w = wb[name]
    out.append(f'=== {name} {cell} = {w[cell].value!r} ; max_row={w.max_row} ===')
    # dump rows around
    rr = int(cell[1:])
    for r in range(max(1, rr - 3), min(w.max_row, rr + 2) + 1):
        vals = []
        for c in range(1, 8):
            v = w.cell(r, c).value
            if v is not None:
                vals.append(f'{get_column_letter(c)}{r}={v!r}')
        if vals:
            out.append('  ' + ' | '.join(vals))

# style templates
def sty(ws, coord):
    c = ws[coord]
    f = c.font; fill = c.fill; al = c.alignment; b = c.border
    return (f'font(name={f.name},sz={f.sz},b={f.b},color={f.color.rgb if f.color else None}) '
            f'fill({fill.fgColor.rgb if fill and fill.fgColor else None},{fill.patternType}) '
            f'numfmt={c.number_format!r} align(h={al.horizontal},v={al.vertical},wrap={al.wrap_text}) '
            f'border(l={b.left.style},r={b.right.style},t={b.top.style},b={b.bottom.style})')

ws = wb['报价汇总']
out.append('=== 报价汇总 styles ===')
for coord in ['A27', 'A28', 'B28', 'A29', 'B29', 'C29', 'D29', 'E29', 'A37', 'B37', 'A38', 'B38', 'A40', 'A41', 'B42', 'A52', 'B53']:
    out.append(f'  {coord}: {sty(ws, coord)}')
out.append('  col widths: ' + repr({k: round(v.width, 1) for k, v in ws.column_dimensions.items() if v.width}))
out.append('  row heights 27-40: ' + repr({r: ws.row_dimensions[r].height for r in range(27, 41) if r in ws.row_dimensions and ws.row_dimensions[r].height}))

ws = wb['制费明细']
out.append('=== 制费明细 styles ===')
for coord in ['A1', 'A3', 'B3', 'A6', 'A7', 'B7', 'A8', 'B8', 'G8', 'C15', 'G15', 'A62', 'H62', 'A64', 'A75', 'A91', 'H91', 'A93']:
    out.append(f'  {coord}: {sty(ws, coord)}')
out.append('  col widths: ' + repr({k: round(v.width, 1) for k, v in ws.column_dimensions.items() if v.width}))
out.append('  merges: ' + repr([str(m) for m in ws.merged_cells.ranges]))

ws = wb['物料明细']
out.append('=== 物料明细 styles ===')
for coord in ['A2', 'B2', 'A3', 'B3', 'C3', 'E3', 'F3', 'G3', 'H3', 'I3', 'A4', 'H4']:
    out.append(f'  {coord}: {sty(ws, coord)}')
out.append('  col widths: ' + repr({k: round(v.width, 1) for k, v in ws.column_dimensions.items() if v.width}))
out.append('  freeze: ' + repr(ws.freeze_panes))

ws = wb['参数']
out.append('=== 参数 styles ===')
for coord in ['A1', 'A2', 'B2', 'A3', 'B3', 'C3', 'A8', 'A9', 'B10', 'A34', 'A41', 'A43', 'B43']:
    out.append(f'  {coord}: {sty(ws, coord)}')
out.append('  col widths: ' + repr({k: round(v.width, 1) for k, v in ws.column_dimensions.items() if v.width}))
out.append('  merges: ' + repr([str(m) for m in ws.merged_cells.ranges]))
out.append('  DVs:')
for dv in ws.data_validations.dataValidation:
    out.append(f'    {dv.sqref} formula1={dv.formula1!r}')

ws = wb['使用说明']
out.append('=== 使用说明 full ===')
for r in range(1, ws.max_row + 1):
    vals = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v is not None:
            vals.append(f'{get_column_letter(c)}{r}={v!r}')
    if vals:
        out.append('  ' + ' | '.join(vals))
out.append('  col widths: ' + repr({k: round(v.width, 1) for k, v in ws.column_dimensions.items() if v.width}))
out.append('  A1 style: ' + sty(ws, 'A1'))
out.append('  A3 style: ' + sty(ws, 'A3'))

with open('/workspace/recon2.txt', 'w') as f:
    f.write('\n'.join(out))
print('OK', len(out))
